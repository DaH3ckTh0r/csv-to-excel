import io
import sys

import pytest
from openpyxl import load_workbook

from csv2xlsx.cli import UsageError, main, parse_args


class TestParseArgs:
    def test_defaults(self):
        a = parse_args(["data/report.csv"])
        assert (a.input, a.output, a.sheet) == ("data/report.csv", "report.xlsx", "Sheet1")
        assert (a.types, a.header, a.table, a.quiet) == (True, True, False, False)

    def test_stdin_output_name(self):
        assert parse_args(["-"]).output == "output.xlsx"

    def test_all_options(self):
        argv = ["in.csv", "out.xlsx", "-d", ";", "--sheet", "Ventas"]
        a = parse_args([*argv, "--no-types", "--no-header", "--table", "-q"])
        assert (a.output, a.delimiter, a.sheet) == ("out.xlsx", ";", "Ventas")
        assert (a.types, a.header, a.table, a.quiet) == (False, False, True, True)

    @pytest.mark.parametrize("d", ["tab", "TAB", "\\t"])
    def test_tab_aliases(self, d):
        assert parse_args(["in.csv", "-d", d]).delimiter == "\t"

    def test_missing_input(self):
        with pytest.raises(UsageError):
            parse_args([])

    def test_bad_delimiter(self):
        with pytest.raises(UsageError, match="single character"):
            parse_args(["in.csv", "-d", ";;"])

    def test_unknown_option(self):
        with pytest.raises(UsageError):
            parse_args(["in.csv", "--bogus"])


class TestMain:
    def test_converts_file(self, tmp_path, monkeypatch, capsys):
        monkeypatch.chdir(tmp_path)
        (tmp_path / "data.csv").write_text("a,b\n1,2\n")
        assert main(["data.csv"]) == 0
        assert "2 rows × 2 cols" in capsys.readouterr().out
        assert load_workbook(tmp_path / "data.xlsx").active["A2"].value == 1

    def test_reads_stdin(self, tmp_path, monkeypatch):
        monkeypatch.setattr(sys, "stdin", io.TextIOWrapper(io.BytesIO(b"x;y\n1;2\n")))
        out = tmp_path / "out.xlsx"
        assert main(["-", str(out), "-q"]) == 0
        assert load_workbook(out).active["B2"].value == 2

    def test_missing_file(self, capsys):
        assert main(["nope.csv"]) == 1
        assert "Cannot read file" in capsys.readouterr().err

    def test_usage_error_exit_code(self, capsys):
        assert main(["in.csv", "-d", ";;"]) == 2
        assert "single character" in capsys.readouterr().err

    def test_help(self, capsys):
        assert main(["--help"]) == 0
        assert "examples:" in capsys.readouterr().out

    def test_version(self, capsys):
        assert main(["--version"]) == 0
        assert "1.0.0" in capsys.readouterr().out

    def test_quiet_prints_nothing(self, tmp_path, monkeypatch, capsys):
        monkeypatch.chdir(tmp_path)
        (tmp_path / "d.csv").write_text("a\n1\n")
        assert main(["d.csv", "-q"]) == 0
        assert capsys.readouterr().out == ""
