import pytest
from openpyxl import load_workbook

from csv2xlsx.converter import ConvertOptions, build_workbook, convert_file, decode

SAMPLE = (
    'id,name,notes,amount,zip,active\r\n1,"Smith, John","line1\nline2",12.5,00123,true\r\n'
    '2,"Say ""hi""",,-3,90210,FALSE\r\n3,,x\r\n'
)


def rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


class TestDecode:
    def test_strips_utf8_bom(self):
        assert decode("﻿a,b".encode()) == "a,b"

    def test_utf16_le_with_bom(self):
        assert decode(b"\xff\xfe" + "á,b".encode("utf-16-le")) == "á,b"

    def test_utf16_be_with_bom(self):
        assert decode(b"\xfe\xff" + "a,b".encode("utf-16-be")) == "a,b"

    def test_falls_back_to_cp1252(self):
        assert decode("café".encode("cp1252")) == "café"

    def test_explicit_encoding(self):
        assert decode(b"\xe9", "latin-1") == "é"


class TestBuildWorkbook:
    def test_type_detection(self):
        wb, stats = build_workbook(SAMPLE)
        assert (stats.rows, stats.columns, stats.delimiter) == (4, 6, ",")
        data = rows(wb["Sheet1"])
        assert data[0] == ["id", "name", "notes", "amount", "zip", "active"]
        assert data[1] == [1, "Smith, John", "line1\nline2", 12.5, "00123", True]
        assert data[2][:2] == [2, 'Say "hi"']
        assert data[2][5] is False

    def test_no_types(self):
        wb, _ = build_workbook("a\n42", ConvertOptions(types=False))
        assert rows(wb.active)[1] == ["42"]

    def test_header_not_coerced(self):
        wb, _ = build_workbook("2024,true\n1,2")
        assert rows(wb.active)[0] == ["2024", "true"]

    def test_no_header_coerces_first_row(self):
        wb, _ = build_workbook("2024,true\n1,2", ConvertOptions(header=False))
        assert rows(wb.active)[0] == [2024, True]

    def test_header_formatting(self):
        ws = build_workbook(SAMPLE)[0].active
        assert ws["A1"].font.bold
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref == "A1:F4"

    def test_table_with_unique_headers(self):
        ws = build_workbook("id,id,\n1,2,3", ConvertOptions(table=True))[0].active
        assert "Data" in ws.tables
        assert rows(ws)[0] == ["id", "id_2", "Column3"]

    def test_sanitized_sheet_name(self):
        wb, _ = build_workbook("a", ConvertOptions(sheet="Ventas/2026"))
        assert wb.active.title == "Ventas_2026"

    def test_detects_semicolon(self):
        assert build_workbook("a;b\n1;2")[1].delimiter == ";"

    def test_pads_short_rows(self):
        assert build_workbook("a,b,c\n1")[1].columns == 3

    def test_warns_unterminated_quote(self):
        assert len(build_workbook('a,"b')[1].warnings) == 1

    def test_strips_illegal_characters(self):
        wb, _ = build_workbook("a\nx\x01y")
        assert rows(wb.active)[1] == ["xy"]

    def test_too_many_columns_raises(self):
        with pytest.raises(ValueError, match="Too many"):
            build_workbook(",".join(["x"] * 16_385))

    def test_empty_input_raises(self):
        with pytest.raises(ValueError, match="empty"):
            build_workbook("\n\n")


def test_convert_file_writes_readable_xlsx(tmp_path):
    src = tmp_path / "in.csv"
    out = tmp_path / "out.xlsx"
    src.write_text(SAMPLE, encoding="utf-8", newline="")
    stats = convert_file(src, out)
    assert stats.rows == 4
    assert rows(load_workbook(out).active)[1][1] == "Smith, John"
