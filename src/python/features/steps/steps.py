"""Step definitions for features/convert.feature (shared with the TypeScript project)."""

import contextlib
import io
import os
import shlex
from pathlib import Path

from behave import given, then, when
from openpyxl import load_workbook

from csv2xlsx.cli import main


def _sheet(context, name):
    path = Path(context.dir) / context.last_output
    if path not in context.workbooks:
        context.workbooks[path] = load_workbook(path)
    wb = context.workbooks[path]
    assert name in wb.sheetnames, f'Sheet "{name}" not found in {wb.sheetnames}'
    return wb[name]


@given("a working directory")
def step_working_dir(context):
    assert os.path.isdir(context.dir)


@given('a CSV file "{name}" containing:')
def step_csv_file(context, name):
    (Path(context.dir) / name).write_text(context.text + "\n", encoding="utf-8")


@when('I run csv2xlsx with "{arg_line}"')
def step_run(context, arg_line):
    args = shlex.split(arg_line)
    positionals = [
        a for i, a in enumerate(args) if not a.startswith("-") and not (i and args[i - 1].startswith("-"))
    ]
    if len(positionals) > 1:
        context.last_output = positionals[1]
    else:
        context.last_output = Path(positionals[0]).with_suffix(".xlsx").name

    err, out = io.StringIO(), io.StringIO()
    cwd = os.getcwd()
    try:
        os.chdir(context.dir)
        with contextlib.redirect_stderr(err), contextlib.redirect_stdout(out):
            context.exit_code = main([*args, "--quiet"])
    finally:
        os.chdir(cwd)
    context.stderr = err.getvalue()


@then("the command succeeds")
def step_succeeds(context):
    assert context.exit_code == 0, f"exit {context.exit_code}, stderr: {context.stderr}"


@then("the command fails with exit code {code:d}")
def step_fails(context, code):
    assert context.exit_code == code, f"expected {code}, got {context.exit_code}"


@then('the error output mentions "{text}"')
def step_error_mentions(context, text):
    assert text in context.stderr, context.stderr


@then('the file "{name}" exists')
def step_file_exists(context, name):
    assert (Path(context.dir) / name).is_file(), f"{name} was not created"


@then('sheet "{name}" has {rows:d} rows and {cols:d} columns')
def step_dimensions(context, name, rows, cols):
    ws = _sheet(context, name)
    assert (ws.max_row, ws.max_column) == (rows, cols), (ws.max_row, ws.max_column)


@then('cell "{ref}" of sheet "{name}" is the text "{expected}"')
def step_cell_text(context, ref, name, expected):
    expected = expected.replace("\\n", "\n").replace('\\"', '"')
    assert _sheet(context, name)[ref].value == expected, repr(_sheet(context, name)[ref].value)


@then('cell "{ref}" of sheet "{name}" is the number {expected:g}')
def step_cell_number(context, ref, name, expected):
    value = _sheet(context, name)[ref].value
    assert isinstance(value, (int, float)) and not isinstance(value, bool), repr(value)
    assert value == expected, repr(value)


@then('cell "{ref}" of sheet "{name}" is the boolean {expected:w}')
def step_cell_bool(context, ref, name, expected):
    assert _sheet(context, name)[ref].value is (expected == "true")


@then('sheet "{name}" contains an Excel table')
def step_has_table(context, name):
    assert len(_sheet(context, name).tables) > 0, "No table found"
